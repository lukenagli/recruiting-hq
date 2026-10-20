"""Each user's writing templates: their cover letter Word file and their email templates.

Cover letter: the user uploads the Word file of a letter in the Yale SOM format (date, recipient block,
"Dear ...:", body paragraphs, "Sincerely,", name). New letters are written into a copy of that file, so the
fonts, margins and spacing are exactly theirs. Without one, the Yale SOM Career Development template is used.

Email templates: the Yale SOM Finance Club templates (Meeting #1 deck) are the defaults; each user can edit them.
Parts in [square brackets] are filled in for each email; everything else is kept word for word.
"""
import io
import json
import re
from pathlib import Path

from .db import connect, init_db

DEFAULT_COVER = Path(__file__).parent / "assets" / "yale_som_cover_letter_template.docx"

EMAIL_KINDS = {
    "outreach": {
        "label": "Coffee chat outreach",
        "help": "The first email asking a banker or 2Y for a coffee chat, with three time windows.",
        "default": (
            "Subject Line: [Student First & Last Name] / [Banker First & Last Name] - Coffee Chat\n\n"
            "Hi [Banker Name],\n\n"
            "Hope you're doing well! It was nice meeting you at the [event]. I'd like to learn more about your path to "
            "Investment Banking and your recruiting experience. Do you have a moment to chat sometime this week? Here is my "
            "availability for the next few days:\n\n"
            "• [Day, date]: [time window] ET\n"
            "• [Day, date]: [time window] ET\n"
            "• [Day, date]: [time window] ET\n\n"
            "Please let me know what time is convenient for you and I will send over a calendar invite. Happy to chat over "
            "Zoom/phone or in-person.\n\n"
            "Best,\n[Student Name]\nMBA Candidate | Yale School of Management '[class year]\n[Mobile Number] | [email]"
        ),
    },
    "calendar_invite": {
        "label": "Calendar invite",
        "help": "The note in the calendar invite once they confirm a time. Attach your resume.",
        "default": (
            "Subject Line: [Student First & Last Name] / [Banker or 2Y First & Last Name] - Coffee Chat\n\n"
            "[Banker or 2Y Name],\n\n"
            "Thank you again for your time. Please find attached a copy of my resume and Zoom link [if applicable].\n\n"
            "I look forward to speaking with you!\n\n"
            "[Zoom Meeting Details, if applicable]\n\n"
            "Best,\n[Student Name]\nMBA Candidate | Yale School of Management '[class year]\n[Mobile Number] | [email]"
        ),
    },
    "post_chat": {
        "label": "Post-coffee chat thank-you",
        "help": "Sent 4 to 24 hours after the chat. The app drafts this one for you after you log a chat.",
        "default": (
            "[First name],\n\nI just want to thank you again for meeting with me [when].\n\n"
            "I enjoyed hearing more about what led you to [Bank] and [something specific about their path or work there]. "
            "I also liked learning about [a second topic from the conversation].\n\n"
            "Thank you again for your time and insight. I really enjoyed our conversation and hope we can stay in touch "
            "throughout the process."
        ),
    },
    "slots_expired": {
        "label": "Time slots expired follow-up",
        "help": "Sent only after every time you offered has passed, with a new set of times. Paste the version you use.",
        "default": "",
    },
}
MAX_EMAIL = 4000


def email_templates(uid: int) -> dict:
    """{kind: text} with the club defaults for anything the user has not set."""
    init_db()
    with connect(read_only=True) as conn:
        row = conn.execute("SELECT email_templates, email_template FROM user_profiles WHERE user_id = ?", (uid,)).fetchone()
    saved = {}
    if row:
        try:
            saved = json.loads(row["email_templates"] or "{}")
        except ValueError:
            saved = {}
        if "post_chat" not in saved and (row["email_template"] or "").strip():
            saved["post_chat"] = row["email_template"]  # from before there were several templates
    return {k: saved.get(k, v["default"]) for k, v in EMAIL_KINDS.items()}


def is_default(kind: str, text: str) -> bool:
    return " ".join(text.split()) == " ".join(EMAIL_KINDS[kind]["default"].split())


def save_email_templates(uid: int, updates: dict) -> dict:
    current = email_templates(uid)
    for kind, text in updates.items():
        if kind in EMAIL_KINDS:
            current[kind] = str(text).replace("\r\n", "\n").strip()[:MAX_EMAIL]
    init_db()
    with connect() as conn:
        conn.execute("INSERT OR IGNORE INTO user_profiles (user_id) VALUES (?)", (uid,))
        conn.execute("UPDATE user_profiles SET email_templates = ?, email_template = '' WHERE user_id = ?", (json.dumps(current), uid))
    return current


def email_options() -> list[dict]:
    return [{"kind": k, "label": v["label"], "help": v["help"], "default": v["default"]} for k, v in EMAIL_KINDS.items()]


# ---------------- Cover letter Word template ----------------

DATE = re.compile(r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+(\d{4}|20XX)", re.I)
CLOSINGS = ("sincerely", "best regards", "regards", "respectfully", "kind regards", "warm regards", "best,", "thank you,")


def layout(doc) -> dict | None:
    """Where the date, recipient block, salutation, body, closing and signature are in a letter, or None."""
    paras = doc.paragraphs
    text = [p.text.strip() for p in paras]
    sal = next((i for i, t in enumerate(text) if t.lower().startswith("dear ")), None)
    if sal is None:
        return None
    close = next((i for i in range(sal + 1, len(paras)) if text[i].lower().startswith(CLOSINGS)), None)
    if close is None:
        return None
    date = next((i for i in range(sal) if DATE.search(text[i])), None)
    recipient = [i for i in range((date + 1) if date is not None else 0, sal) if text[i]]
    body = [i for i in range(sal + 1, close) if text[i]]
    signature = next((i for i in range(len(paras) - 1, close, -1) if text[i]), None)
    if len(body) < 2:
        return None
    return {"date": date, "recipient": recipient, "salutation": sal, "body": body, "close": close, "signature": signature}


def check_cover_template(data: bytes) -> tuple[str, str]:
    """(plain text, problem) for an uploaded Word file. problem is '' when it can be used."""
    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception:
        return "", "That doesn't look like a Word (.docx) file."
    if layout(doc) is None:
        return "", ("Couldn't find the letter's parts. Upload a cover letter in the Yale SOM format: a \"Dear ...:\" line, "
                    "your body paragraphs, then \"Sincerely,\" and your name.")
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip()), ""


def save_cover_template(uid: int, filename: str, data: bytes) -> dict:
    text, problem = check_cover_template(data)
    if problem:
        return {"ok": False, "reason": problem}
    init_db()
    with connect() as conn:
        conn.execute("INSERT OR IGNORE INTO user_profiles (user_id) VALUES (?)", (uid,))
        conn.execute("UPDATE user_profiles SET cover_template = ?, cover_template_name = ?, example_letter = ? WHERE user_id = ?",
                     (data, filename[:200], text[:8000], uid))
    return {"ok": True, **cover_template_info(uid)}


def cover_template_info(uid: int) -> dict:
    init_db()
    with connect(read_only=True) as conn:
        row = conn.execute("SELECT cover_template_name, length(cover_template) AS size FROM user_profiles WHERE user_id = ?", (uid,)).fetchone()
    has = bool(row and row["size"])
    return {"has_template": has, "filename": (row["cover_template_name"] if has else "") or ""}


def cover_template_bytes(uid: int) -> bytes:
    init_db()
    with connect(read_only=True) as conn:
        row = conn.execute("SELECT cover_template FROM user_profiles WHERE user_id = ?", (uid,)).fetchone()
    data = row["cover_template"] if row else None
    return bytes(data) if data else DEFAULT_COVER.read_bytes()


def remove_cover_template(uid: int) -> dict:
    init_db()
    with connect() as conn:
        conn.execute("UPDATE user_profiles SET cover_template = NULL, cover_template_name = '' WHERE user_id = ?", (uid,))
    return cover_template_info(uid)
