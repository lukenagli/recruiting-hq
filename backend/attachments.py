"""Files dropped into the assistant: a bank's team card, a saved email, a screenshot of an invite.

Text is pulled out here (PDF, Word, plain text, .eml) so the assistant reads it like a pasted message. A scanned
PDF or an image has no text to pull out, so it is handed to the model to read directly. The file is not kept.
"""
import io
import re
from email import policy
from email.parser import BytesParser

from pydantic_ai import BinaryContent

MAX_BYTES = 10_000_000
MAX_CHARS = 50_000
IMAGES = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp", "gif": "image/gif"}


class AttachmentError(Exception):
    """A message that is safe to show the person who attached the file."""


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) > 60:
        raise AttachmentError(f"That PDF has {len(reader.pages)} pages; attach one under 60 pages.")
    return "\n\n".join((p.extract_text() or "").strip() for p in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # team cards and rosters are often laid out as tables
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _eml(data: bytes) -> str:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    head = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        text = re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style).*?</\1>", " ", text))
    return f"{head}\n\n{text}"


def _msg(data: bytes) -> str:
    """An email dragged out of desktop Outlook (.msg, a Microsoft compound file of MAPI properties)."""
    import struct
    from datetime import datetime, timedelta

    import olefile

    ole = olefile.OleFileIO(io.BytesIO(data))

    def prop(code: str) -> str:
        for kind, codec in (("001F", "utf-16-le"), ("001E", "cp1252")):
            name = f"__substg1.0_{code}{kind}"
            if ole.exists(name):
                return ole.openstream(name).read().decode(codec, errors="replace").rstrip("\x00").strip()
        return ""

    sent = ""
    if ole.exists("__properties_version1.0"):  # fixed-size properties after a 32-byte header, 16 bytes each
        raw = ole.openstream("__properties_version1.0").read()
        for i in range(32, len(raw) - 15, 16):
            tag, _, value = struct.unpack_from("<IIQ", raw, i)
            if tag == 0x00390040 and value:  # PR_CLIENT_SUBMIT_TIME, a Windows FILETIME
                sent = (datetime(1601, 1, 1) + timedelta(microseconds=value // 10)).strftime("%Y-%m-%d %H:%M UTC")
    sender, address = prop("0C1A"), prop("0C1F") or prop("5D01")
    head = [("From", f"{sender} <{address}>" if sender and "@" in address else sender or address),
            ("To", prop("0E04")), ("Cc", prop("0E03")), ("Date", sent), ("Subject", prop("0037"))]
    body = prop("1000")
    if not body and ole.exists("__substg1.0_10130102"):  # only an HTML body was saved
        html = ole.openstream("__substg1.0_10130102").read().decode("utf-8", errors="replace")
        body = re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style).*?</\1>", " ", html))
    ole.close()
    return "\n".join(f"{k}: {v}" for k, v in head if v) + f"\n\n{body}"


def read(filename: str, data: bytes, outlook: bool = False) -> tuple[str, BinaryContent | None]:
    """(text, file for the model to read itself). Exactly one of the two is filled in.
    `outlook` turns on reading .msg emails dragged straight out of desktop Outlook."""
    if len(data) > MAX_BYTES:
        raise AttachmentError("That file is over 10 MB.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == "pdf":
            text = _pdf(data)
            if len(text.strip()) < 200:  # a scan or a picture of a page: let the model read it
                return "", BinaryContent(data=data, media_type="application/pdf")
        elif ext == "docx":
            text = _docx(data)
        elif ext == "eml":
            text = _eml(data)
        elif ext in ("txt", "md", "csv"):
            text = data.decode("utf-8", errors="replace")
        elif ext in IMAGES:
            return "", BinaryContent(data=data, media_type=IMAGES[ext])
        elif ext == "msg" and outlook:
            text = _msg(data)
        elif ext == "msg":
            raise AttachmentError("Outlook .msg files can't be read. Open the email and paste its text, or save it as .eml.")
        else:
            raise AttachmentError("Attach a PDF, Word file, image, text file or saved email (.eml).")
    except AttachmentError:
        raise
    except Exception:
        raise AttachmentError("Couldn't read that file. Try saving it again or pasting its text.")
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        raise AttachmentError("That file has no text in it.")
    return text[:MAX_CHARS], None
