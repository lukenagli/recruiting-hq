"""Load the question bank workbooks and the application tracker into SQLite.

Run with:  python -m backend.importer
Safe to re-run: a question already in the database is skipped, so each run reports only what is new.
"""
import hashlib
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

from . import config, normalize as nz
from .db import connect, init_db

LIVE_FILE = "Question Bank.xlsx"
MASTER_FILE = "000. Yale SOM Finance Club Master Question Bank_vShared.xlsx"
IB_2025_FILE = "2025 - IB Question Bank.xlsx"
TECHNICAL_FILE = "Technical Question Bank.xlsx"


def _rows(path: Path, sheet: str) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return []
        return [list(r) for r in wb[sheet].iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    finally:
        wb.close()


def _table(rows: list[list], required: set[str]) -> list[dict]:
    """Find the header row (the first row holding every required column) and return the rows below it as dicts."""
    for i, row in enumerate(rows):
        header = [nz.clean(c) for c in row]
        if required <= set(header):
            return [
                {name: (r[j] if j < len(r) else None) for j, name in enumerate(header) if name}
                for r in rows[i + 1:]
            ]
    return []


def _is_template_row(rec: dict) -> bool:
    """The live log has a hint row ("e.g. Goldman, JPM...") and an EXAMPLE row under the header."""
    bank, question = nz.clean(rec.get("Bank")), nz.clean(rec.get("Question"))
    return bank.lower().startswith("e.g.") or question.upper().startswith("EXAMPLE")


def _question(source: str, rec: dict, **fixed) -> dict | None:
    question = nz.clean(rec.get("Question"))
    if not question or _is_template_row(rec):
        return None
    return {
        "source": source,
        "year": nz.optional(rec.get("Year")),
        "qtype": nz.qtype(rec.get("Type") or rec.get("Type of Question")),
        "bank": nz.bank(rec.get("Bank")),
        "office": nz.office(rec.get("Office") or rec.get("Location")),
        "grp": nz.optional(rec.get("Group")),
        "round": nz.round_(rec.get("Round")),
        "interviewer": nz.optional(rec.get("Interviewer")),
        "topic": nz.optional(rec.get("Topic")),
        "question": question,
        "answer_notes": nz.optional(rec.get("Answer / Notes")),
        **fixed,
    }


def read_live(path: Path) -> dict[str, list[dict]]:
    out = {}
    for source, sheet in (("live_1y", "Interview Question Log - 1Ys"), ("live_2y", "Interview Question Log - 2Ys")):
        recs = _table(_rows(path, sheet), {"Bank", "Round", "Question"})
        out[source] = [q for q in (_question(source, r, year="2026") for r in recs) if q]
    return out


def read_master(path: Path) -> list[dict]:
    recs = _table(_rows(path, "Master Question Bank"), {"Bank", "Round", "Question"})
    return [q for q in (_question("master", r) for r in recs) if q]


def read_ib_2025(path: Path) -> list[dict]:
    recs = _table(_rows(path, "IB Question Bank (2025)"), {"Bank", "Round", "Question"})
    return [q for q in (_question("ib_2025", r, year="2025") for r in recs) if q]


def read_technical(path: Path) -> list[dict]:
    """Topic heading rows, then numbered questions. An unnumbered long row continues the question above it."""
    out, topic = [], ""
    for row in _rows(path, "Question Bank"):
        cells = [nz.clean(c) for c in row if c not in (None, "")]
        if len(cells) >= 2 and cells[0].isdigit():
            out.append(_question("technical", {"Type": "T", "Topic": topic, "Question": cells[1]}))
        elif len(cells) == 1 and len(cells[0]) > 40 and out:
            out[-1]["question"] += " " + cells[0]
        elif len(cells) == 1:
            topic = cells[0]
    return [q for q in out if q]


def _iso(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return nz.clean(value)


def read_applications(path: Path) -> list[dict]:
    recs = _table(_rows(path, "Application Tracker - 1Ys"), {"Bank", "Official Deadline"})
    return [
        {
            "bank": nz.bank(r.get("Bank")),
            "official_deadline": _iso(r.get("Official Deadline")),
            "recommended_date": _iso(r.get("Recommended Submission Date")),
            "link": nz.clean(r.get("Application Link")),
            "cover_letter": nz.clean(r.get("Cover Letter?")),
            "other_info": nz.clean(r.get("Other Information")),
        }
        for r in recs
        if nz.clean(r.get("Bank"))
    ]


def _dedupe_key(q: dict) -> str:
    parts = [q["source"], q["bank"], q["grp"], q["round"], q["interviewer"], q["question"]]
    return hashlib.sha256("|".join(p.lower() for p in parts).encode("utf-8")).hexdigest()


def _save_questions(conn, source: str, path: Path, questions: list[dict], now: str) -> dict:
    new = 0
    for q in questions:
        cur = conn.execute(
            """INSERT OR IGNORE INTO questions
               (source, year, qtype, bank, office, grp, round, interviewer, topic, question,
                answer_notes, dedupe_key, first_seen)
               VALUES (:source, :year, :qtype, :bank, :office, :grp, :round, :interviewer, :topic,
                       :question, :answer_notes, :dedupe_key, :first_seen)""",
            {**q, "dedupe_key": _dedupe_key(q), "first_seen": now},
        )
        new += cur.rowcount
    return _log(conn, source, path, len(questions), new, now)


def _save_applications(conn, path: Path, apps: list[dict], now: str) -> dict:
    for a in apps:
        conn.execute(
            """INSERT INTO applications
               (bank, official_deadline, recommended_date, link, cover_letter, other_info, updated_at)
               VALUES (:bank, :official_deadline, :recommended_date, :link, :cover_letter, :other_info, :now)
               ON CONFLICT(bank) DO UPDATE SET
                 official_deadline = excluded.official_deadline,
                 recommended_date  = excluded.recommended_date,
                 link              = excluded.link,
                 cover_letter      = excluded.cover_letter,
                 other_info        = excluded.other_info,
                 updated_at        = excluded.updated_at""",
            {**a, "now": now},
        )
    return _log(conn, "applications", path, len(apps), 0, now)


def _log(conn, source: str, path: Path, rows_read: int, rows_new: int, now: str) -> dict:
    mtime = datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    conn.execute(
        "INSERT INTO imports (source, file, file_mtime, rows_read, rows_new, ran_at) VALUES (?, ?, ?, ?, ?, ?)",
        (source, path.name, mtime, rows_read, rows_new, now),
    )
    return {"source": source, "file": path.name, "file_mtime": mtime, "rows_read": rows_read, "rows_new": rows_new}


def import_live_file(path: Path) -> dict:
    """Import just the live Finance Club workbook (question logs and application tracker) from a given path."""
    init_db()
    now = datetime.now().isoformat(timespec="seconds")
    conn = connect()
    try:
        with conn:
            results = [_save_questions(conn, source, path, questions, now) for source, questions in read_live(path).items()]
            results.append(_save_applications(conn, path, read_applications(path), now))
        total = conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0]
    finally:
        conn.close()
    return {"ran_at": now, "results": results, "total_questions": total}


def run_import() -> dict:
    """Import every workbook found in the data folder. Missing files are reported, not fatal."""
    folder = config.banks_dir()
    init_db()
    now = datetime.now().isoformat(timespec="seconds")
    results, missing = [], []
    conn = connect()
    try:
        with conn:
            live = folder / LIVE_FILE
            if live.is_file():
                for source, questions in read_live(live).items():
                    results.append(_save_questions(conn, source, live, questions, now))
                results.append(_save_applications(conn, live, read_applications(live), now))
            else:
                missing.append(LIVE_FILE)
            for source, name, reader in (
                ("master", MASTER_FILE, read_master),
                ("ib_2025", IB_2025_FILE, read_ib_2025),
                ("technical", TECHNICAL_FILE, read_technical),
            ):
                path = folder / name
                if path.is_file():
                    results.append(_save_questions(conn, source, path, reader(path), now))
                else:
                    missing.append(name)
        total = conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0]
    finally:
        conn.close()
    return {"folder": str(folder), "ran_at": now, "results": results, "missing": missing, "total_questions": total}


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    summary = run_import()
    print(f"Data folder: {summary['folder']}")
    for r in summary["results"]:
        print(f"  {r['source']:<13} read {r['rows_read']:>5}  new {r['rows_new']:>5}  ({r['file']}, saved {r['file_mtime']})")
    for name in summary["missing"]:
        print(f"  MISSING: {name}")
    print(f"Total questions in database: {summary['total_questions']}")
