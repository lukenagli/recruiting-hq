"""Each user's profile: who they are, what they are recruiting for, and the templates their writing follows."""
import io
from datetime import datetime

from .db import connect, init_db

FIELDS = ("resume_text", "story", "example_letter", "contact_line", "school_line", "school", "level", "class_year",
          "target_groups", "email_template", "email_rules", "why_group", "highlights", "projects")
MAX_CHARS = {"resume_text": 15_000, "story": 6_000, "example_letter": 8_000, "contact_line": 300, "school_line": 200,
             "school": 120, "level": 20, "class_year": 20, "target_groups": 400, "email_template": 4_000, "email_rules": 2_000,
             "why_group": 3_000, "highlights": 3_000, "projects": 8_000}


def get(uid: int) -> dict:
    init_db()
    with connect(read_only=True) as conn:
        row = conn.execute("SELECT * FROM user_profiles WHERE user_id = ?", (uid,)).fetchone()
    return {f: (row[f] if row else "") or "" for f in FIELDS} | {
        "updated_at": row["updated_at"] if row else "", "onboarded_at": (row["onboarded_at"] if row else "") or ""}


def save(uid: int, fields: dict) -> dict:
    clean = {f: str(fields.get(f, "")).strip()[: MAX_CHARS[f]] for f in FIELDS if f in fields}
    init_db()
    with connect() as conn:
        conn.execute("INSERT OR IGNORE INTO user_profiles (user_id) VALUES (?)", (uid,))
        for field, value in clean.items():
            conn.execute(f"UPDATE user_profiles SET {field} = ? WHERE user_id = ?", (value, uid))
        conn.execute("UPDATE user_profiles SET updated_at = ? WHERE user_id = ?", (datetime.now().isoformat(timespec="seconds"), uid))
    return get(uid)


def background(prof: dict) -> str:
    """Everything the user wrote about themselves, labelled, for the cover letter writer and the interview coach."""
    parts = [
        ("Target groups, in order of preference", prof.get("target_groups", "")),
        ("Why those groups", prof.get("why_group", "")),
        ("What they want to highlight", prof.get("highlights", "")),
        ("Notable projects and experiences to draw on", prof.get("projects", "")),
        ("Other notes", prof.get("story", "")),
    ]
    return (chr(10) * 2).join(f"{label}:{chr(10)}{text.strip()}" for label, text in parts if (text or "").strip())


def finish_setup(uid: int) -> None:
    init_db()
    with connect() as conn:
        conn.execute("INSERT OR IGNORE INTO user_profiles (user_id) VALUES (?)", (uid,))
        conn.execute("UPDATE user_profiles SET onboarded_at = ? WHERE user_id = ?", (datetime.now().isoformat(timespec="seconds"), uid))


def text_from_upload(filename: str, data: bytes) -> str:
    """Pull plain text out of an uploaded resume (.docx, .pdf or .txt)."""
    name = filename.lower()
    if name.endswith(".docx"):
        import docx

        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        return "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="replace")
    raise ValueError("Upload a .docx, .pdf or .txt file.")
